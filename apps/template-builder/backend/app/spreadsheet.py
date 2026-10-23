"""Read spreadsheets (.xlsx / .xlsm / .csv / .tsv) as rows of display text.

Cells come back the way a person sees them in the spreadsheet: numbers with
their number format applied (decimals, thousands separators, currency, per
cent) and dates formatted, because that text is what goes into a template.
Bold and italic text in Excel cells comes back too (see ``_cell_styles``), so a
cell that is partly bold stays that way in the template.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from pathlib import PurePath

MAX_ROWS = 2000
MAX_COLUMNS = 60
MAX_SHEETS = 20


class SpreadsheetError(Exception):
    """The file cannot be read as a spreadsheet; the message is safe to show."""


_SECTION_COLOR = re.compile(r"\[(?:Black|Blue|Cyan|Green|Magenta|Red|White|Yellow|Color\d+)\]", re.I)
_CONDITION = re.compile(r"\[[<>=][^\]]*\]")
_CURRENCY = re.compile(r"\[\$([^\-\]]*)(?:-[0-9A-Fa-f]+)?\]")
_QUOTED = re.compile(r'"([^"]*)"')
_DATE_TOKENS = re.compile(r"yyyy|yy|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|AM/PM|A/P", re.I)


def _literal_parts(section: str) -> tuple[str, str, str]:
    """Split a number-format section into (prefix, digit pattern, suffix)."""
    section = _CURRENCY.sub(lambda m: m.group(1), section)
    # Protect quoted text and escaped characters from being read as placeholders.
    literals: list[str] = []

    def stash(text: str) -> str:
        literals.append(text)
        return f"\x00{len(literals) - 1}\x00"

    section = _QUOTED.sub(lambda m: stash(m.group(1)), section)
    section = re.sub(r"\\(.)", lambda m: stash(m.group(1)), section)
    section = re.sub(r"_.", " ", section)      # "_)" leaves a space the width of ")"
    section = re.sub(r"\*.", "", section)      # "*x" repeats x to fill the cell
    match = re.search(r"[0#?][0#?,.]*(?:E[+-]?0+)?", section)
    if not match:
        prefix, pattern, suffix = section, "", ""
    else:
        prefix, pattern, suffix = section[:match.start()], match.group(0), section[match.end():]

    def restore(text: str) -> str:
        return re.sub(r"\x00(\d+)\x00", lambda m: literals[int(m.group(1))], text)

    return restore(prefix), pattern, restore(suffix)


def format_number(value: float, number_format: str) -> str:
    """Apply an Excel number format to a number (the common cases)."""
    fmt = _CONDITION.sub("", _SECTION_COLOR.sub("", number_format or "General"))
    sections = fmt.split(";")
    section = sections[0]
    negative = value < 0
    if negative and len(sections) > 1 and sections[1].strip():
        section, value = sections[1], -value      # the negative section carries its own sign / brackets
        negative = False
    elif value == 0 and len(sections) > 2 and sections[2].strip():
        section = sections[2]

    if section.strip().lower() in ("general", "@", ""):
        if float(value).is_integer() and abs(value) < 1e15:
            return str(int(value))
        return f"{value:.10g}"

    prefix, pattern, suffix = _literal_parts(section)
    sign = "-" if negative else ""
    if not pattern:
        return (prefix + suffix).strip() or f"{value:.10g}"
    if "%" in prefix + suffix:
        value *= 100
    if "E" in pattern.upper():
        decimals = len(pattern.split(".")[1].split("E")[0]) if "." in pattern else 0
        return f"{prefix}{value:.{decimals}E}{suffix}".strip()
    integer_part, _, decimal_part = pattern.partition(".")
    # trailing commas scale by thousands; a comma between digits groups them
    scale = len(integer_part) - len(integer_part.rstrip(","))
    value /= 1000 ** scale
    grouping = "," in integer_part.rstrip(",")
    decimals = sum(ch in "0#?" for ch in decimal_part)
    text = f"{abs(value):,.{decimals}f}" if grouping else f"{abs(value):.{decimals}f}"
    if decimal_part and "0" not in decimal_part:      # "#.##" drops trailing zeros
        text = text.rstrip("0").rstrip(".")
    return f"{sign}{prefix}{text}{suffix}".strip()


def format_date(value: dt.date | dt.time | dt.datetime, number_format: str) -> str:
    fmt = _CONDITION.sub("", _SECTION_COLOR.sub("", number_format or "")).split(";")[0]
    fmt = _CURRENCY.sub("", fmt)
    has_date = isinstance(value, (dt.date, dt.datetime))
    has_time = isinstance(value, (dt.time, dt.datetime))
    if not _DATE_TOKENS.search(fmt) or fmt.strip().lower() == "general":
        if isinstance(value, dt.datetime) and (value.hour or value.minute or value.second):
            return value.strftime("%Y-%m-%d %H:%M")
        return value.strftime("%Y-%m-%d") if has_date else value.strftime("%H:%M")
    twelve_hour = bool(re.search(r"AM/PM|A/P", fmt, re.I))
    out: list[str] = []
    position = 0
    previous = ""
    for match in _DATE_TOKENS.finditer(fmt):
        out.append(re.sub(r'\\(.)|"([^"]*)"', lambda m: m.group(1) or m.group(2) or "", fmt[position:match.start()]))
        position = match.end()
        token = match.group(0).lower()
        rest = fmt[match.end():].lstrip(": ").lower()
        if token in ("m", "mm") and (previous in ("h", "hh") or rest.startswith("s")):
            out.append(f"{value.minute:02d}" if has_time else "00")      # minutes, not month
        elif token == "yyyy":
            out.append(f"{value.year:04d}" if has_date else "")
        elif token == "yy":
            out.append(f"{value.year % 100:02d}" if has_date else "")
        elif token == "mmmm":
            out.append(value.strftime("%B") if has_date else "")
        elif token == "mmm":
            out.append(value.strftime("%b") if has_date else "")
        elif token == "mm":
            out.append(f"{value.month:02d}" if has_date else "")
        elif token == "m":
            out.append(str(value.month) if has_date else "")
        elif token == "dddd":
            out.append(value.strftime("%A") if has_date else "")
        elif token == "ddd":
            out.append(value.strftime("%a") if has_date else "")
        elif token == "dd":
            out.append(f"{value.day:02d}" if has_date else "")
        elif token == "d":
            out.append(str(value.day) if has_date else "")
        elif token in ("h", "hh"):
            hour = value.hour if has_time else 0
            if twelve_hour:
                hour = hour % 12 or 12
            out.append(f"{hour:02d}" if token == "hh" else str(hour))
        elif token in ("s", "ss"):
            second = value.second if has_time else 0
            out.append(f"{second:02d}" if token == "ss" else str(second))
        elif token in ("am/pm", "a/p"):
            late = has_time and value.hour >= 12
            out.append(("PM" if late else "AM") if token == "am/pm" else ("P" if late else "A"))
        previous = token
    out.append(re.sub(r'\\(.)|"([^"]*)"', lambda m: m.group(1) or m.group(2) or "", fmt[position:]))
    return "".join(out).strip()


def cell_text(value, number_format: str = "General") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return format_date(value, number_format)
    if isinstance(value, dt.timedelta):
        total = int(value.total_seconds())
        return f"{total // 3600}:{total % 3600 // 60:02d}:{total % 60:02d}"
    if isinstance(value, (int, float)):
        try:
            return format_number(float(value), number_format)
        except (ValueError, IndexError):
            return f"{value:.10g}"
    return str(value)


def _trim(rows: list[list[str]]) -> list[list[str]]:
    """Drop empty rows at the end and empty columns on the right; make rows equal length."""
    while rows and not any(cell.strip() for cell in rows[-1]):
        rows.pop()
    width = max((max((i + 1 for i, cell in enumerate(row) if cell.strip()), default=0) for row in rows), default=0)
    return [(row + [""] * width)[:width] for row in rows]


def _sheet(name: str, rows: list[list[str]], truncated: bool, styles: dict | None = None) -> dict:
    rows = _trim(rows)
    width = len(rows[0]) if rows else 0
    kept = {key: spans for key, spans in (styles or {}).items()
            if int(key.split(":")[0]) < len(rows) and int(key.split(":")[1]) < width}
    return {"name": name, "rows": rows, "rowCount": len(rows),
            "columnCount": width, "truncated": truncated, "styles": kept}


def _flags(font, inherited: tuple[bool, bool] = (False, False)) -> tuple[bool, bool]:
    """(bold, italic) of an Excel font; unset parts inherit."""
    if font is None:
        return inherited
    bold = inherited[0] if getattr(font, "b", None) is None else bool(font.b)
    italic = inherited[1] if getattr(font, "i", None) is None else bool(font.i)
    return bold, italic


def _cell_styles(cell, text: str) -> list[dict] | None:
    """Bold / italic stretches of a cell, or None when it is all plain.

    ``start`` and ``end`` count only the characters that are not white space,
    so the stretches stay put when the text is wrapped onto more lines.
    """
    base = _flags(getattr(cell, "font", None))
    value = cell.value
    parts: list[tuple[str, tuple[bool, bool]]]
    if type(value).__name__ == "CellRichText":
        parts = [(part, base) if isinstance(part, str) else (part.text, _flags(part.font, base)) for part in value]
    else:
        parts = [(text, base)]
    spans: list[dict] = []
    index = 0
    for part, (bold, italic) in parts:
        count = sum(not ch.isspace() for ch in part or "")
        if count and (bold or italic):
            last = spans[-1] if spans else None
            if last and last["end"] == index and last.get("bold", False) == bold and last.get("italic", False) == italic:
                last["end"] += count
            else:
                span = {"start": index, "end": index + count}
                if bold:
                    span["bold"] = True
                if italic:
                    span["italic"] = True
                spans.append(span)
        index += count
    return spans or None


def _read_workbook(data: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:   # pragma: no cover - dependency is in requirements.txt
        raise SpreadsheetError("Excel support is not installed on the server (pip install openpyxl).") from exc
    try:
        # data_only: the cached result of formulas, not the formulas themselves.
        # Not read_only: that mode drops bold / italic inside a cell.
        workbook = load_workbook(io.BytesIO(data), data_only=True, rich_text=True)
    except Exception as exc:
        raise SpreadsheetError("This file could not be opened as an Excel workbook.") from exc
    sheets: list[dict] = []
    try:
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            if getattr(worksheet, "sheet_state", "visible") != "visible":
                continue
            rows: list[list[str]] = []
            styles: dict[str, list[dict]] = {}
            truncated = False
            for row in worksheet.iter_rows():
                if len(rows) >= MAX_ROWS:
                    truncated = True
                    break
                texts: list[str] = []
                for column, cell in enumerate(row[:MAX_COLUMNS]):
                    value = cell.value
                    if type(value).__name__ == "CellRichText":
                        text = str(value)
                    else:
                        text = cell_text(value, getattr(cell, "number_format", "General") or "General")
                    texts.append(text)
                    spans = _cell_styles(cell, text) if text.strip() else None
                    if spans:
                        styles[f"{len(rows)}:{column}"] = spans
                rows.append(texts)
            sheets.append(_sheet(worksheet.title, rows, truncated, styles))
    finally:
        workbook.close()
    return sheets


def _read_delimited(data: bytes, name: str, delimiter: str | None) -> list[dict]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    if delimiter is None:
        try:
            dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
    else:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows: list[list[str]] = []
    truncated = False
    for row in reader:
        if len(rows) >= MAX_ROWS:
            truncated = True
            break
        rows.append([cell.strip() for cell in row[:MAX_COLUMNS]])
    return [_sheet(name, rows, truncated)]


def read_spreadsheet(data: bytes, filename: str) -> list[dict]:
    """Sheets of ``{name, rows, rowCount, columnCount, truncated, styles}``; rows hold display text.

    ``styles`` maps ``"row:column"`` to the cell's bold / italic stretches (Excel only).
    """
    suffix = PurePath(filename or "").suffix.lower()
    stem = PurePath(filename or "Sheet1").stem or "Sheet1"
    if data[:4] == b"PK\x03\x04" or suffix in (".xlsx", ".xlsm"):
        sheets = _read_workbook(data)
    elif data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or suffix == ".xls":
        raise SpreadsheetError("Old .xls files are not supported. Save the workbook as .xlsx and try again.")
    elif suffix == ".tsv":
        sheets = _read_delimited(data, stem, "\t")
    else:
        sheets = _read_delimited(data, stem, None)
    sheets = [sheet for sheet in sheets if sheet["rowCount"]]
    if not sheets:
        raise SpreadsheetError("This file has no data.")
    return sheets
