from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.extraction import detect_tables_in_file, extract_template
from app.spreadsheet import SpreadsheetError, cell_text, format_number, read_spreadsheet
from app.storage import LocalStorage

from .conftest import upload


@pytest.fixture(scope="module")
def storage(tmp_path_factory) -> LocalStorage:
    return LocalStorage(tmp_path_factory.mktemp("storage"))


def cells(page: dict, table: dict) -> list[list[str]]:
    """Text of each cell, found the way the editor finds it: by position."""
    cols, rows = table["columns"], table["rows"]
    grid = [["" for _ in cols[1:]] for _ in rows[1:]]
    for el in page["elements"]:
        if el["type"] != "text":
            continue
        cx, cy = el["x"] + el["width"] / 2, el["y"] + el["height"] / 2
        if cols[0] <= cx <= cols[-1] and rows[0] <= cy <= rows[-1]:
            c = max(i for i, edge in enumerate(cols[:-1]) if edge <= cx)
            r = max(i for i, edge in enumerate(rows[:-1]) if edge <= cy)
            grid[r][c] = (grid[r][c] + " " + el["content"]).strip()
    return grid


# ------------------------------------------------------------ detection
def test_ruled_table_is_detected(samples: Path, storage):
    page = extract_template(samples / "invoice.pdf", "t-invoice", "invoice", storage)["pages"][0]
    items = next(t for t in page["tables"] if t["source"] == "ruled")
    grid = cells(page, items)
    assert (len(grid), len(grid[0])) == (6, 4) and items["headerRows"] == 1
    assert grid[0] == ["Description", "Qty", "Unit price", "Amount"]
    assert grid[3] == ["Website UI design (per page)", "8", "£320.00", "£2,560.00"]
    # columns of numbers are right-aligned, the description column is not
    by_text = {el.get("content"): el for el in page["elements"] if el["type"] == "text"}
    assert by_text["£320.00"]["align"] == "right" and by_text["Qty"]["align"] == "right"
    assert by_text["Illustration pack"]["align"] == "left"


def test_small_aligned_tables_on_the_invoice(samples: Path, storage):
    page = extract_template(samples / "invoice.pdf", "t-invoice2", "invoice", storage)["pages"][0]
    grids = [cells(page, t) for t in page["tables"]]
    assert [["Subtotal", "£8,620.00"], ["VAT (20%)", "£1,724.00"], ["Total due", "£10,344.00"]] in grids
    assert any(g[0] == ["Invoice number", "INV-2026-0142"] and len(g) == 4 for g in grids)
    # the address block beside the invoice details is not pulled into a table
    address = next(el for el in page["elements"] if el.get("content", "").startswith("Attn:"))
    assert address["content"].count("\n") == 3


def test_table_with_only_horizontal_rules(samples: Path, storage):
    template = extract_template(samples / "multipage.pdf", "t-multi", "multipage", storage)
    page = template["pages"][3]
    assert len(page["tables"]) == 1
    table = page["tables"][0]
    grid = cells(page, table)
    # (rules between the rows: the row-ruled strategy finds it first, with the same grid)
    assert table["source"] in ("aligned", "ruled") and table["headerRows"] == 1
    assert grid[0] == ["Region", "Sites", "Visits", "Change"] and len(grid) == 8
    assert grid[-1] == ["Islands", "60", "3,130", "-13.4%"]
    # the grid spans the rules, and rows do not overlap
    assert table["x"] == pytest.approx(40, abs=0.5) and table["width"] == pytest.approx(515.3, abs=1)
    assert table["rows"] == sorted(table["rows"]) and table["columns"] == sorted(table["columns"])


def test_no_tables_in_prose_charts_or_photo_grids(samples: Path, storage):
    multipage = extract_template(samples / "multipage.pdf", "t-multi2", "multipage", storage)
    assert [len(p["tables"]) for p in multipage["pages"]] == [0, 0, 0, 1, 0, 0, 0, 0]   # p5 is a bar chart
    for name in ("brochure", "certificate"):   # photo captions, three columns of prose, signatures
        template = extract_template(samples / f"{name}.pdf", f"t-{name}", name, storage)
        assert all(not p["tables"] for p in template["pages"]), name


def test_browser_made_table(samples: Path, storage):
    page = extract_template(samples / "newsletter.pdf", "t-news", "newsletter", storage)["pages"][0]
    assert len(page["tables"]) == 1
    grid = cells(page, page["tables"][0])
    assert grid[0] == ["Departs", "From", "Fare"] and grid[1] == ["06:40", "Eastquay", "£3.20"] and len(grid) == 6
    by_text = {el.get("content"): el for el in page["elements"] if el["type"] == "text"}
    assert by_text["06:40"]["align"] == "left"      # same-width values under a left-aligned heading
    assert by_text["£2.60"]["align"] == "right"


def test_table_running_over_pages(samples: Path, storage):
    template = extract_template(samples / "statement.pdf", "t-statement", "statement", storage)
    pages = template["pages"]
    assert [len(p["tables"]) for p in pages] == [2, 1, 1]
    transactions = [next(t for t in p["tables"] if len(t["columns"]) == 8) for p in pages]
    grids = [cells(p, t) for p, t in zip(pages, transactions)]
    assert [len(g) - 1 for g in grids] == [9, 12, 12]                     # body rows per page
    assert all(t["headerRows"] == 1 and t["source"] == "ruled" for t in transactions)
    # every page's piece has the same headings: that is what lets the editor chain them
    assert all(g[0] == ["Date", "Value date", "Particulars", "Cheque", "Debit", "Credit", "Balance"] for g in grids)
    # text wrapped inside one cell stays one element; rows are never merged together
    wrapped = [el for el in pages[1]["elements"] if el["type"] == "text" and "\n" in el["content"]]
    assert wrapped and all(el["content"].count("\n") == 1 for el in wrapped)
    balances = [el for el in pages[1]["elements"] if el["type"] == "text"
                and transactions[1]["columns"][6] < el["x"] and el["y"] > transactions[1]["rows"][1]
                and el["y"] < transactions[1]["rows"][-1]]
    assert len(balances) == 12 and all(el["align"] == "right" for el in balances)


def test_detection_without_full_extraction(samples: Path):
    found = detect_tables_in_file(samples / "invoice.pdf")
    assert len(found) == 1 and sorted(t["source"] for t in found[0]) == ["aligned", "aligned", "ruled"]


# ---------------------------------------------------------- spreadsheets
@pytest.mark.parametrize("value,fmt,expected", [
    (1800, "General", "1800"),
    (3.14159, "General", "3.14159"),
    (1800, "#,##0.00", "1,800.00"),
    (1800, "£#,##0.00", "£1,800.00"),
    (1800, "[$£-809]#,##0.00", "£1,800.00"),
    (-1234.5, "#,##0.00", "-1,234.50"),
    (-1234.5, "$#,##0.00", "-$1,234.50"),
    (-1234.5, "#,##0.00_);(#,##0.00)", "(1,234.50)"),
    (0.256, "0%", "26%"),
    (0.256, "0.0%", "25.6%"),
    (12, "0", "12"),
    (12.5, "0.00", "12.50"),
    (1234567, "#,##0", "1,234,567"),
    (7, '0" kg"', "7 kg"),
    (1250, '_-* #,##0.00_-;-* #,##0.00_-;_-* "-"??_-;_-@_-', "1,250.00"),
])
def test_number_formats(value, fmt, expected):
    assert format_number(value, fmt) == expected


def test_date_and_other_cells():
    day = dt.datetime(2026, 10, 1, 14, 5)
    assert cell_text(day, "dd/mm/yyyy") == "01/10/2026"
    assert cell_text(day, "d mmm yyyy") == "1 Oct 2026"
    assert cell_text(day, "mmmm d, yyyy") == "October 1, 2026"
    assert cell_text(day, "yyyy-mm-dd hh:mm") == "2026-10-01 14:05"
    assert cell_text(day, "h:mm AM/PM") == "2:05 PM"
    assert cell_text(dt.date(2026, 10, 1), "General") == "2026-10-01"
    assert cell_text(None) == "" and cell_text(True) == "TRUE" and cell_text("  text ") == "  text "


def workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(["Description", "Qty", "Unit price", "Amount"])
    for row in [("Consulting", 3, 450, 1350), ("Travel", 1, 89.5, 89.5)]:
        ws.append(row)
    for cell in ws["C"][1:] + ws["D"][1:]:
        cell.number_format = "£#,##0.00"
    ws["F2"] = dt.date(2026, 10, 1)
    ws["F2"].number_format = "dd/mm/yyyy"
    ws.append([])                     # trailing empty rows are dropped
    other = wb.create_sheet("Notes")
    other.append(["only", "one", "row"])
    wb.create_sheet("Empty")          # sheets without data are dropped
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_read_workbook():
    sheets = read_spreadsheet(workbook_bytes(), "items.xlsx")
    assert [s["name"] for s in sheets] == ["Items", "Notes"]
    items = sheets[0]
    assert items["rows"][0][:4] == ["Description", "Qty", "Unit price", "Amount"]
    assert items["rows"][1] == ["Consulting", "3", "£450.00", "£1,350.00", "", "01/10/2026"]
    assert items["rows"][2][:4] == ["Travel", "1", "£89.50", "£89.50"]
    assert (items["rowCount"], items["columnCount"], items["truncated"]) == (3, 6, False)


def test_read_bold_and_italic():
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.append(["Item", "Notes"])
    sheet.append(["Pen", CellRichText(["Colour: ", TextBlock(InlineFont(b=True), "Blue ink"), " pack of 2"])])
    sheet.append(["Total", "plain"])
    sheet["A3"].font = Font(bold=True, italic=True)
    out = io.BytesIO()
    book.save(out)
    read = read_spreadsheet(out.getvalue(), "styled.xlsx")[0]
    assert read["rows"][1][1] == "Colour: Blue ink pack of 2"
    # counted over the characters that are not spaces: "Colour:" is 7, "Blueink" the next 7
    assert read["styles"] == {
        "1:1": [{"start": 7, "end": 14, "bold": True}],
        "2:0": [{"start": 0, "end": 5, "bold": True, "italic": True}],
    }


def test_read_csv_and_errors():
    sheets = read_spreadsheet("name;amount\nA;1,5\nB;2\n".encode(), "data.csv")
    assert sheets[0]["rows"] == [["name", "amount"], ["A", "1,5"], ["B", "2"]]
    assert read_spreadsheet(b"a\tb\n1\t2\n", "data.tsv")[0]["rows"] == [["a", "b"], ["1", "2"]]
    with pytest.raises(SpreadsheetError, match="no data"):
        read_spreadsheet(b"", "empty.csv")
    with pytest.raises(SpreadsheetError, match="xlsx"):
        read_spreadsheet(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 old excel", "legacy.xls")
    with pytest.raises(SpreadsheetError, match="could not be opened"):
        read_spreadsheet(b"PK\x03\x04 not really a zip", "broken.xlsx")


# ------------------------------------------------------------------ API
def test_parse_endpoint(client):
    res = client.post("/api/spreadsheets/parse", files={"file": ("items.xlsx", workbook_bytes(), "application/octet-stream")})
    assert res.status_code == 200
    body = res.json()
    assert body["filename"] == "items.xlsx" and body["sheets"][0]["rows"][1][2] == "£450.00"
    bad = client.post("/api/spreadsheets/parse", files={"file": ("x.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "application/vnd.ms-excel")})
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "invalid_spreadsheet"


def test_tables_are_stored_and_backfilled(client, samples):
    template = upload(client, samples / "invoice.pdf").json()["template"]
    assert len(template["pages"][0]["tables"]) == 3
    saved = client.get(f"/api/templates/{template['id']}").json()
    assert saved["pages"][0]["tables"] == template["pages"][0]["tables"]

    # a template imported before table detection existed has NULL tables: they are found on first open
    from app.db import Page, SessionLocal
    with SessionLocal() as db:
        db.query(Page).filter(Page.template_id == template["id"]).update({Page.tables: None})
        db.commit()
    reopened = client.get(f"/api/templates/{template['id']}").json()
    assert [t["source"] for t in reopened["pages"][0]["tables"]] == [t["source"] for t in template["pages"][0]["tables"]]
    with SessionLocal() as db:
        assert db.query(Page).filter(Page.template_id == template["id"]).one().tables is not None

    # tables survive saving edits and duplication
    client.put(f"/api/templates/{template['id']}", json={"pages": [reopened["pages"][0]]})
    copy = client.post(f"/api/templates/{template['id']}/duplicate").json()
    assert len(client.get(f"/api/templates/{copy['id']}").json()["pages"][0]["tables"]) == 3
